"""AS9102 First Article Inspection Report workbook: Forms 1, 2 and 3.

Field numbering follows AS9102 Rev C. Fields the drawing cannot answer
(serial number, PO, signatures, results) are left blank for the inspector.
Confirm the layout against the form revision your customer requires.
"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from src import normalize

THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill("solid", fgColor="D9D9D9")
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")
BOLD = Font(bold=True, size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
NOTICE = ("Layout follows AS9102 Rev C field numbering. Generated from the drawing; "
          "every entry must be verified by the inspector before release.")

TYPE_LABEL = {"linear": "Dimension", "diameter": "Diameter", "radius": "Radius",
              "angle": "Angle", "thread": "Thread", "gdt": "GD&T", "note": "Note"}


# ---------- helpers ----------

def requirement(entry: dict) -> str:
    """Requirement text as the drawing states it, without repeating a frame twice."""
    v = normalize.clean(entry.get("value"))
    tol = normalize.clean(entry.get("tolerance"))
    if not tol:
        return v
    nv, nt = normalize.canonical(v), normalize.canonical(tol)
    if nv and nv in nt:
        return tol
    if nt in nv:
        return v
    if entry.get("type") == "gdt" and normalize.gdt_class(v) and not normalize.numbers(v):
        return tol  # value is just the characteristic name, the frame says it all
    return f"{v} {tol}"


def limits(entry: dict) -> tuple[float, float, float] | None:
    """(nominal, lower, upper) for simple bilateral or limit tolerances, else None."""
    if entry.get("type") == "gdt":
        return None
    value = re.sub(r"^\s*\d+\s*[xX]\s*", "", normalize.clean(entry.get("value")))  # drop 4X
    nums = normalize.numbers(value)
    tol = normalize.clean(entry.get("tolerance"))
    if len(nums) != 1 or not tol:
        return None
    n = nums[0]
    m = re.fullmatch(r"±\s*(\d*\.?\d+)", tol)
    if m:
        t = float(m.group(1))
        return n, round(n - t, 9), round(n + t, 9)
    m = re.fullmatch(r"\+\s*(\d*\.?\d+)\s*/\s*-\s*(\d*\.?\d+)", tol)
    if m:
        return n, round(n - float(m.group(2)), 9), round(n + float(m.group(1)), 9)
    return None


def _header(ws, title: str, fields: list[tuple[str, str]], width: int):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=12)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    ws["A2"] = NOTICE
    ws["A2"].fill = NOTE_FILL
    ws["A2"].alignment = WRAP
    r = 3
    for i in range(0, len(fields), 2):
        for j, (label, value) in enumerate(fields[i:i + 2]):
            c = 1 + j * 4
            ws.cell(row=r, column=c, value=label).font = BOLD
            ws.merge_cells(start_row=r, start_column=c + 1, end_row=r, end_column=c + 3)
            ws.cell(row=r, column=c + 1, value=value or "")
        r += 1
    return r + 1


def _table(ws, row: int, columns: list[tuple[str, int]], rows: list[list]):
    for i, (name, w) in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=i, value=name)
        cell.font, cell.fill, cell.border, cell.alignment = BOLD, HDR_FILL, BOX, WRAP
        ws.column_dimensions[cell.column_letter].width = max(ws.column_dimensions[cell.column_letter].width or 0, w)
    for r, values in enumerate(rows, start=row + 1):
        for c, val in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.border, cell.alignment = BOX, WRAP
    return row + len(rows) + 2


# ---------- forms ----------

def form1(wb, info: dict):
    ws = wb.active
    ws.title = "Form 1"
    fields = [("1. Part Number", info.get("part_number")), ("2. Part Name", info.get("part_name")),
              ("3. Serial Number", info.get("serial_number")), ("4. FAI Report Number", info.get("fai_report_number")),
              ("5. Part Revision Level", info.get("revision")), ("6. Drawing Number", info.get("drawing_number")),
              ("7. Drawing Revision Level", info.get("revision")), ("8. Additional Changes", ""),
              ("9. Manufacturing Process Reference", ""), ("10. Organization Name", info.get("organization")),
              ("11. Supplier Code", ""), ("12. Purchase Order Number", info.get("order_number")),
              ("13. Detail FAI / Assembly FAI", "Detail FAI"), ("14. Full FAI / Partial FAI", "Full FAI")]
    row = _header(ws, "AS9102 First Article Inspection - Form 1: Part Number Accountability", fields, 8)
    row = _table(ws, row, [("15. Part Number", 18), ("16. Part Name", 22), ("17. Part Serial Number", 18), ("18. FAI Report Number", 18)],
                 [["", "", "", ""] for _ in range(4)])
    sign = [("19. Does FAIR contain a documented nonconformance?", ""), ("20. FAI Complete / FAI Not Complete", ""),
            ("21. Signature", ""), ("22. Date", ""), ("23. Reviewed By", ""), ("24. Date", ""),
            ("25. Customer Approval", ""), ("26. Date", "")]
    for i, (label, _) in enumerate(sign):
        ws.cell(row=row + i, column=1, value=label).font = BOLD
        ws.merge_cells(start_row=row + i, start_column=2, end_row=row + i, end_column=4)


def form2(wb, info: dict):
    ws = wb.create_sheet("Form 2")
    fields = [("1. Part Number", info.get("part_number")), ("2. Part Name", info.get("part_name")),
              ("3. Serial Number", info.get("serial_number")), ("4. FAI Report Number", info.get("fai_report_number"))]
    row = _header(ws, "AS9102 First Article Inspection - Form 2: Product Accountability (Raw Material, Specifications, Special Processes, Functional Testing)", fields, 9)
    specs = list(info.get("specifications") or [])
    if info.get("material") and not any(s.get("type") == "material" for s in specs if isinstance(s, dict)):
        specs.insert(0, {"name": info["material"], "spec": "", "type": "material"})
    rows = []
    for s in specs:
        if not isinstance(s, dict):
            continue
        is_test = s.get("type") == "test"
        rows.append([s.get("name") or "", s.get("spec") or "", "", "", "", "",
                     s.get("spec") if is_test else "", "", s.get("type") or ""])
    rows += [[""] * 9 for _ in range(max(3, 8 - len(rows)))]
    _table(ws, row, [("5. Material or Process Name", 28), ("6. Specification Number", 22), ("7. Code", 10),
                     ("8. Special Process Supplier Code", 18), ("9. Customer Approval Verification (Yes/No/NA)", 18),
                     ("10. Certificate of Conformance Number", 18), ("11. Functional Test Procedure Number", 18),
                     ("12. Acceptance Report Number", 16), ("13. Comments", 24)], rows)


def form3(wb, info: dict, entries: list[dict]):
    ws = wb.create_sheet("Form 3")
    fields = [("1. Part Number", info.get("part_number")), ("2. Part Name", info.get("part_name")),
              ("3. Serial Number", info.get("serial_number")), ("4. FAI Report Number", info.get("fai_report_number"))]
    row = _header(ws, "AS9102 First Article Inspection - Form 3: Characteristic Accountability, Verification and Compatibility Evaluation", fields, 8)
    rows = []
    for e in sorted((e for e in entries if e.get("balloon")), key=lambda e: e["balloon"]):
        kind = TYPE_LABEL.get(e.get("type"), "")
        if e.get("type") == "gdt":
            kind = f"GD&T {normalize.gdt_class(e.get('tolerance')) or normalize.gdt_class(e.get('value')) or ''}".strip()
        comments = [kind]
        if e.get("bbox_source") == "estimate":
            comments.append("balloon position approximate")
        rows.append([e["balloon"], e.get("reference_location") or "", "KC" if e.get("key") else "",
                     requirement(e), "", "", "", "; ".join(c for c in comments if c)])
    _table(ws, row, [("5. Char No.", 8), ("6. Reference Location", 16), ("7. Characteristic Designator", 14),
                     ("8. Requirement", 38), ("9. Results", 16), ("10. Designed / Qualified Tooling", 20),
                     ("11. Nonconformance Number", 18), ("12. Additional Data / Comments", 30)], rows)
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def plan_sheet(wb, entries: list[dict]):
    """Not an AS9102 form: limits worked out from the drawing, to speed up inspection."""
    ws = wb.create_sheet("Inspection Plan")
    ws["A1"] = "Inspection plan (working sheet, not part of AS9102). Limits are calculated from simple ± or +/- tolerances only."
    ws["A1"].font = BOLD
    rows = []
    for e in sorted((e for e in entries if e.get("balloon")), key=lambda e: e["balloon"]):
        lim = limits(e)
        nums = [round(v, 6) for v in lim] if lim else ["", "", ""]
        rows.append([e["balloon"], e.get("reference_location") or "", requirement(e), *nums,
                     e.get("bbox_source") or ""])
    _table(ws, 3, [("Char No.", 8), ("Location", 14), ("Requirement", 38), ("Nominal", 12),
                   ("Lower limit", 12), ("Upper limit", 12), ("Position found by", 16)], rows)


def build_workbook(entries: list[dict], info: dict, out_path: str | Path) -> Path:
    wb = Workbook()
    form1(wb, info)
    form2(wb, info)
    form3(wb, info, entries)
    plan_sheet(wb, entries)
    out = Path(out_path)
    wb.save(out)
    return out
