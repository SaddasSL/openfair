import json

from openpyxl import load_workbook

from eval.score import score
from src.pipeline import run_pipeline
from tests.conftest import CALLOUTS, SAMPLES, fake_extractor_for


def test_black_drawing_end_to_end(black_drawing, tmp_path):
    res = run_pipeline(str(black_drawing), "PN-1", "Test part", out_dir=str(tmp_path),
                       extractor=fake_extractor_for(CALLOUTS), title_extractor=lambda p: {"revision": "B"})
    data = json.loads(open(res["json"], encoding="utf-8").read())
    chars = data["characteristics"]
    assert res["characteristics"] == 4
    assert {c["value"] for c in chars} == {"Ø10.5", "Ø1.5", "25.40", "R5"}
    assert all(c["bbox_source"] == "pdf-text" for c in chars)
    zones = {c["value"]: c["reference_location"] for c in chars}
    assert zones == {"Ø10.5": "Sh 1 / B4", "Ø1.5": "Sh 1 / B4", "25.40": "Sh 1 / B2", "R5": "Sh 1 / A1"}
    wb = load_workbook(res["workbook"])
    assert "Form 2" in wb.sheetnames
    assert res["ballooned_pdf"].endswith(".pdf")


def test_multi_page_nist_ctc5_no_crash(tmp_path):
    """Two page drawing, coloured PMI, model returns nothing useful: still completes."""
    res = run_pipeline(str(SAMPLES / "nist_ctc5.pdf"), out_dir=str(tmp_path),
                       extractor=lambda p: [{"type": "linear", "value": "5.000", "tolerance": "±.008 (ST)", "bbox": [50, 50, 120, 30]}],
                       title_extractor=lambda p: (_ for _ in ()).throw(RuntimeError("no key")))
    assert res["pages"] == 2
    assert res["warnings"] and "title block" in res["warnings"][0]
    data = json.loads(open(res["json"], encoding="utf-8").read())
    assert {c["page"] for c in data["characteristics"]} == {1, 2}


def test_from_json_rebuild(black_drawing, tmp_path):
    first = run_pipeline(str(black_drawing), out_dir=str(tmp_path / "a"), extractor=fake_extractor_for(CALLOUTS),
                         title_extractor=lambda p: {})
    again = run_pipeline(str(black_drawing), out_dir=str(tmp_path / "b"), from_json=first["json"])
    assert again["characteristics"] == first["characteristics"]


def test_scoring_against_answer_key():
    key = json.loads((SAMPLES.parents[1] / "eval" / "answer_keys" / "nist_ctc3.json").read_text(encoding="utf-8"))
    perfect = [dict(c) for c in key["characteristics"]]
    s = score(perfect, key)
    assert s["recall"] == 1 and s["precision"] == 1
    reformatted = [{**c, "value": c["value"].replace("Ø.", "Ø0.")} for c in perfect]
    assert score(reformatted, key)["recall"] == 1
    s = score(perfect[2:], key)
    assert s["found_required"] == s["required"] - 2
