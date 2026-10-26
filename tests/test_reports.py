from openpyxl import load_workbook

from src.reports.fair import build_workbook, limits, requirement


def test_requirement_text():
    assert requirement({"type": "diameter", "value": "Ø10.5", "tolerance": "±0.1"}) == "Ø10.5 ±0.1"
    assert requirement({"type": "gdt", "value": "position", "tolerance": "⌖ Ø.05 | A | B"}) == "⌖ Ø.05 | A | B"
    assert requirement({"type": "gdt", "value": "⌖ Ø.05 | A", "tolerance": "⌖ Ø.05 | A"}) == "⌖ Ø.05 | A"


def test_limits():
    assert limits({"type": "linear", "value": "25.40", "tolerance": "±0.10"}) == (25.4, 25.3, 25.5)
    lo = limits({"type": "diameter", "value": "4X Ø.625", "tolerance": "+.002/-.001"})
    assert lo[0] == 0.625 and abs(lo[1] - 0.624) < 1e-9 and abs(lo[2] - 0.627) < 1e-9
    assert limits({"type": "gdt", "value": "position", "tolerance": "⌖ .1 | A"}) is None


def test_workbook_has_three_forms(tmp_path):
    entries = [{"type": "diameter", "value": "Ø10.5", "tolerance": "±0.1", "balloon": 1, "reference_location": "Sh 1 / B4"},
               {"type": "gdt", "value": "flatness", "tolerance": "⏥ .005", "balloon": 2, "key": True},
               {"type": "note", "value": "datum A"}]
    info = {"part_number": "PN-1", "part_name": "Bracket", "revision": "C", "material": "AL 6061-T6",
            "specifications": [{"name": "Anodise", "spec": "MIL-A-8625 Type II", "type": "process"}]}
    out = build_workbook(entries, info, tmp_path / "fair.xlsx")
    wb = load_workbook(out)
    assert wb.sheetnames[:3] == ["Form 1", "Form 2", "Form 3"]
    f1 = [c.value for row in wb["Form 1"].iter_rows() for c in row]
    assert "PN-1" in f1 and "Bracket" in f1 and "26. Date" in f1
    f2 = [c.value for row in wb["Form 2"].iter_rows() for c in row]
    assert "AL 6061-T6" in f2 and "MIL-A-8625 Type II" in f2
    rows = [[c.value for c in row] for row in wb["Form 3"].iter_rows()]
    data = [r for r in rows if isinstance(r[0], int)]
    assert [r[0] for r in data] == [1, 2]
    assert data[0][1] == "Sh 1 / B4" and data[0][3] == "Ø10.5 ±0.1"
    assert data[1][2] == "KC"
